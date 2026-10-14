"""Document parsers: route by extension, convert to markdown, extract image assets.

Every parser returns a markdown string. Structure fidelity (heading levels,
tables, images) is the priority because downstream splitting and retrieval
depend on it.
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

MAX_TABLE_ROWS = 200

IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"}


def parse_document(src, assets_dir: Path, doc_id: str) -> dict:
    """Parse a document. Returns {'markdown': str, 'title': str}."""
    src = Path(src)
    assets_dir = Path(assets_dir)
    assets_dir.mkdir(parents=True, exist_ok=True)
    ext = src.suffix.lower()
    handlers = {
        ".docx": _parse_docx,
        ".xlsx": _parse_xlsx, ".xlsm": _parse_xlsx, ".csv": _parse_csv,
        ".pptx": _parse_pptx,
        ".pdf": _parse_pdf,
        ".txt": _parse_text, ".md": _parse_text,
    }
    if ext in IMAGE_EXTS:
        md = _parse_image(src, assets_dir, doc_id)
    elif ext in handlers:
        md = handlers[ext](src, assets_dir, doc_id)
    else:
        raise ValueError(f"不支持的文档格式: {ext}")
    title = src.stem
    m = re.search(r"^#\s+(.+)$", md, re.M)
    if m:
        title = m.group(1).strip()
    return {"markdown": md, "title": title}


# ---------------- Word ----------------

def _parse_docx(src: Path, assets_dir: Path, doc_id: str) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    from docx.oxml.ns import qn

    doc = Document(str(src))
    out: list[str] = []
    img_counter = {"n": 0}

    def save_image(part) -> str:
        img_counter["n"] += 1
        ext = Path(str(part.partname)).suffix or ".png"
        fname = f"{doc_id}_img{img_counter['n']:02d}{ext}"
        (assets_dir / fname).write_bytes(part.blob)
        return f"![图片{img_counter['n']}](../../assets/{doc_id}/{fname})"

    def para_md(p: Paragraph) -> list[str]:
        lines = []
        text = p.text.strip()
        style = (p.style.name or "") if p.style is not None else ""
        if text:
            m = re.match(r"Heading\s*(\d)", style)
            if m:
                lines.append("#" * min(int(m.group(1)), 6) + " " + text)
            elif style.startswith("List"):
                lines.append("- " + text)
            else:
                lines.append(text)
        for blip in p._element.findall(".//" + qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if not rid:
                continue
            part = doc.part.related_parts.get(rid)
            if part is not None:
                lines.append(save_image(part))
        return lines

    def table_md(t: Table) -> str:
        rows = [[c.text.strip().replace("\n", "<br>") for c in r.cells] for r in t.rows]
        if not rows:
            return ""
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        header, body = rows[0], rows[1:]
        md = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
        md += ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join(md)

    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            out.extend(para_md(Paragraph(child, doc)))
        elif child.tag == qn("w:tbl"):
            t = table_md(Table(child, doc))
            if t:
                out.append(t)
    return "\n\n".join(out)


# ---------------- Excel / CSV ----------------

def _rows_to_md(rows: list[list[str]], truncated: bool) -> str:
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:]
    md = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
    md += ["| " + " | ".join(r) + " |" for r in body]
    if truncated:
        md.append(f"> 注：表格过大，仅转换前 {MAX_TABLE_ROWS} 行。")
    return "\n".join(md)


def _parse_xlsx(src: Path, assets_dir: Path, doc_id: str) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(src), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## {ws.title}")
        rows, truncated = [], False
        for row in ws.iter_rows(values_only=True):
            vals = ["" if v is None else str(v) for v in row]
            if any(v.strip() for v in vals):
                rows.append(vals)
            if len(rows) > MAX_TABLE_ROWS:
                truncated = True
                rows = rows[:MAX_TABLE_ROWS]
                break
        out.append(_rows_to_md(rows, truncated) if rows else "（空工作表）")
    return "\n\n".join(out)


def _parse_csv(src: Path, assets_dir: Path, doc_id: str) -> str:
    import csv

    text = src.read_text(encoding="utf-8", errors="replace")
    rows = [r for r in csv.reader(text.splitlines()) if any(c.strip() for c in r)]
    if not rows:
        return f"# {src.stem}\n\n（空文件）"
    truncated = len(rows) > MAX_TABLE_ROWS
    return _rows_to_md(rows[:MAX_TABLE_ROWS], truncated)


# ---------------- PowerPoint ----------------

def _parse_pptx(src: Path, assets_dir: Path, doc_id: str) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(src))
    out: list[str] = []
    img_n = 0
    for i, slide in enumerate(prs.slides, 1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.text:
            title = slide.shapes.title.text.strip()
        out.append(f"## 第{i}页 {title}".rstrip())
        for shape in slide.shapes:
            if shape is slide.shapes.title:
                continue
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img_n += 1
                ext = shape.image.ext or "png"
                fname = f"{doc_id}_img{img_n:02d}.{ext}"
                (assets_dir / fname).write_bytes(shape.image.blob)
                out.append(f"![图片{img_n}](../../assets/{doc_id}/{fname})")
            elif shape.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                if rows:
                    width = max(len(r) for r in rows)
                    rows = [r + [""] * (width - len(r)) for r in rows]
                    md = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
                    md += ["| " + " | ".join(r) + " |" for r in rows[1:]]
                    out.append("\n".join(md))
            elif shape.has_text_frame:
                is_body = shape.is_placeholder and shape.placeholder_format.idx != 0
                for para in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in para.runs).strip()
                    if t:
                        out.append("- " + t if is_body else t)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                out.append(f"> 备注：{notes}")
    return "\n\n".join(out)


# ---------------- PDF ----------------

def _parse_pdf(src: Path, assets_dir: Path, doc_id: str) -> str:
    import fitz  # pymupdf

    doc = fitz.open(str(src))
    lines: list[tuple[int, float, str]] = []
    sizes: dict[float, int] = {}
    for pno, page in enumerate(doc, 1):
        for block in page.get_text("dict").get("blocks", []):
            if block.get("type") != 0:
                continue
            for line in block.get("lines", []):
                spans = line.get("spans", [])
                text = "".join(s["text"] for s in spans).strip()
                if not text:
                    continue
                size = round(max(s["size"] for s in spans), 1)
                lines.append((pno, size, text))
                sizes[size] = sizes.get(size, 0) + len(text)
    if not lines:
        return f"# {src.stem}\n\n> 扫描版 PDF，未提取到文字内容，待接入 OCR 后处理。"
    body_size = max(sizes, key=sizes.get)
    big_sizes = sorted([s for s in sizes if s > body_size * 1.15], reverse=True)[:2]
    level_of = {s: i + 1 for i, s in enumerate(big_sizes)}

    out: list[str] = []
    cur_page = 0
    for pno, size, text in lines:
        if pno != cur_page:
            out.append(f"<!-- page: {pno} -->")
            cur_page = pno
        if size in level_of:
            out.append("#" * level_of[size] + " " + text)
        else:
            out.append(text)
    return "\n\n".join(out)


# ---------------- images & plain text ----------------

def _parse_image(src: Path, assets_dir: Path, doc_id: str) -> str:
    fname = f"{doc_id}{src.suffix.lower()}"
    shutil.copy2(src, assets_dir / fname)
    return (f"# {src.stem}\n\n![{src.stem}](../../assets/{doc_id}/{fname})\n\n"
            "> 图片资料：文字内容待 OCR 提取，图像描述待视觉模型补充。")


def _parse_text(src: Path, assets_dir: Path, doc_id: str) -> str:
    text = src.read_text(encoding="utf-8", errors="replace").strip()
    if src.suffix.lower() == ".md":
        return text
    return f"# {src.stem}\n\n{text}"
